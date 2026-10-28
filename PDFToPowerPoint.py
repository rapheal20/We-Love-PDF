#This function converts a PDF file to a .pptx
#Coded by Rapheal


import os 
import io
from helperfunctions import *
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
import math
from CommonFunctionsForDocs import *

pdf1 = "Inputs\\emaar1.pdf"

def PDFToPowerPoint(PDFPath):
    extractedElements, firstBox = parse_pdf_to_spans(PDFPath)

    def addPPTXRun(para,text, base, size, color):
        #run is a python-pptx function that represents a portion of text that has same formattin etc
        #para is also a feature of python-pptx, represents one paragraph

        run = para.add_run() #creates new text run inside the PP para
        run.text = text # assigns the text to the run
        family, bold, italic = fontStyle(base) #fontstyle gives the font name
        #assigns the pdf text properties to the run
        run.font.name = family #defines font family
        run.font.size = Pt(max(1, round(size, 1))) #defines font size
        run.font.bold = bold #checks if bold
        run.font.italic = italic #checks if italics

        #above line apply all the properties for the run


        if color:
            run.font.color.rgb = RGBColor(*[int(v) for v in color])
            #makes sure no floats are present

    def PPTXBuilder(spans, firstBox, PDFPath):
        #============Initializing the PPTX============ 
        prs = Presentation()
        # prs.slide_width = Inches(10)
        # prs.slide_height = Inches(7.5)

        #Dimensions in points
        # slideWidth_PTS = 720
        # slideHeight_PTS = 540

        IDENT = (1, 0, 0, 1, 0, 0)
        #this identity matrix reprents no scaleing and no movement
        slide = None #current slide
        tf = None #current textbox, None mean start a new textbox
        para = None #current para
        first_page = True #makes sure slide size is set only once
        x0, y1, pdf_W, pdf_H = 0, 540, 720, 540
        # y1 = top edge, used for flipping, rest are dimensions
        base, fsize, color, leading, mode = None, 12, (0,0,0), 0, 0
        # base = fontname, fsize = size, color = color, leading = for next line, mode = text rendering mode
        ctm = tm = IDENT    
        #ctm is the maatrix of current block
        #tm is th ecurrent text position matrix

        for span in spans:
            t = span.get('type')

            if t == 'pageStart' or span.get('page'):
                #if pagesize is missing US Letter page is assumed
                box = span.get('box') or (0, 0, 612, 792) #page rectangle
                x0, y1 = box[0], box[3]
                pdf_W, pdf_H = box[2] - box[0], box[3] - box[1]
                #using the edges works out its width and height
                if first_page:
                    prs.slide_width = Pt(pdf_W)
                    prs.slide_height = Pt(pdf_H)
                    first_page = False
                    # page size defined only with the 1st page
                    #no need for scaleing as size are set using page size
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                #adds a blank slide, layout 6 is the empty slide
                tf = None
                continue

            if slide is None:
                continue

            if t == 'image':
                try:
                    pic = slide.shapes.add_picture(
                        io.BytesIO(span['data']),
                        left=Pt(span['x'] - x0),
                        top=Pt(y1 - span['y']), #to account for the difference in 0 in y axis
                        width=Pt(span['w']),
                        height=Pt(span['h'] if span.get('h') else None)
                    )

                    # sends the pictures to the back of layers
                    slide.shapes._spTree.remove(pic._element)
                    slide.shapes._spTree.insert(2, pic._element)
                except Exception:
                    #skips that image
                    pass
                continue


            if t == 'blockStart':
                # this is where a new block starts
                # its transformations are stored in ctm
                # tm is reset ie position is reset
                # tf is None, to make a new textbox for new text
                ctm, tm, tf = tuple(span['ctm']), IDENT, None
                color = span['color']
            elif t == 'font':
                base, fsize = span['base'], span['size']
            elif t == 'matrix':
                # tm stores the new positions of the text
                # tf is None, to make a new textbox
                tm = (span['a'], span['b'], span['c'], span['d'], span['x'], span['y'])
                tf = None
            elif t in ('position', 'newline'):
                if t == 'position':
                    dx, dy = span['dx'], span['dy']
                    if span['operator'] == 'TD':
                        leading = -dy
                else:
                    dx, dy = 0 , -leading

                a, b, c, d, e, f = tm
                # account for any scaling or rotation
                tm = (a, b, c, d, e + dx * a + dy * c, f + dx * b + dy * d)
                tf = None
            elif t == 'leading':
                leading = span['value']
            elif t == 'color':
                color = span['rgb']
            elif t == 'renderingMode':
                mode = span['mode']

            elif t == 'text':
                text = span.get('content', '')
                if text == '' or mode in (3, 7):
                    # mode 3 and 7 are empty layes so are skipped
                    continue

                x = tm[4] * ctm[0] + tm[5] * ctm[2] + ctm[4]
                y = tm[4] * ctm[1] + tm[5] * ctm[3] + ctm[5]
                # converts the text position into the page coordintes by combining text matrix with blocks transformation
                mc = tm[2] * ctm[0] + tm[3] * ctm[2]
                md = tm[2] * ctm[1] + tm[3] * ctm[3]
                size = abs(fsize) * math.hypot(mc, md)
                #scales the font size based on vertical scale

                if tf is None:
                    txBox = slide.shapes.add_textbox(
                        Pt(x - x0), Pt(y1 - y - size * 0.9),
                        Pt(max(size, pdf_W - (x - x0))), Pt(size * 1.2))
                    tf = txBox.text_frame
                    tf.word_wrap = True
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    para = tf.paragraphs[0]

                addPPTXRun(para, text, base, size, color)

        PDFName = os.path.basename(PDFPath).split(".")[0]
        outputPath = f"Outputs/{PDFName}.pptx"
        prs.save(outputPath)
        print("done")

    PPTXBuilder(extractedElements, firstBox, PDFPath)



PDFToPowerPoint(pdf1)